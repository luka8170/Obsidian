#!/usr/bin/env python3

from docx import Document

print("="*80)
print("Word 파일 테이블 구조 분석")
print("="*80)

word_file = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/내부 회의 내부 작업/**2. QG 계약서/POSSM/최종 정리/내가 보낸 파일들/word/25.8.6. 송부) FYF AGREEMENT - POSSM(PAN OCEAN)_(as of Apr 30 2025) rev2.docx"

doc = Document(word_file)

print(f"\nTotal tables: {len(doc.tables)}")
print("\n검색: DOK, TPSC, TSBC가 포함된 테이블 찾기...\n")

for table_idx, table in enumerate(doc.tables):
    has_target = False
    target_rows = []
    
    for row_idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        row_text = " | ".join(cells)
        
        if any(keyword in row_text for keyword in ["DOK", "TPSC", "TSBC", "Periodical Bottom", "Tailshaft"]):
            has_target = True
            target_rows.append((row_idx, row_text))
    
    if has_target:
        print(f"\n{'='*80}")
        print(f"Table {table_idx} - Found target surveys!")
        print(f"{'='*80}")
        for row_idx, row_text in target_rows:
            print(f"  Row {row_idx}: {row_text}")

print("\n" + "="*80)
