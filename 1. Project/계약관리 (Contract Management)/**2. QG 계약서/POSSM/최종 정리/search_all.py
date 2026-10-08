#!/usr/bin/env python3

from docx import Document

print("="*80)
print("Word 파일 전체 텍스트 검색")
print("="*80)

word_file = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/내부 회의 내부 작업/**2. QG 계약서/POSSM/최종 정리/내가 보낸 파일들/word/25.8.6. 송부) FYF AGREEMENT - POSSM(PAN OCEAN)_(as of Apr 30 2025) rev2.docx"

doc = Document(word_file)

print("\n[1] Searching in paragraphs...")
annex_c_found = False
context_lines = []

for para_idx, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    
    if "ANNEX C" in text:
        annex_c_found = True
        print(f"\nFound 'ANNEX C' at paragraph {para_idx}: {text}")
    
    if annex_c_found:
        if any(keyword in text for keyword in ["DOK", "TPSC", "TSBC"]):
            print(f"  Paragraph {para_idx}: {text}")
            context_lines.append(text)
        
        # Stop after finding enough context
        if len(context_lines) > 20:
            break

print("\n[2] Searching in tables (showing all cells)...")
for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        cells_text = []
        for cell in row.cells:
            cells_text.append(cell.text.strip())
        
        full_row = " | ".join(cells_text)
        
        if any(keyword in full_row for keyword in ["DOK", "TPSC", "TSBC"]):
            print(f"\nTable {table_idx}, Row {row_idx}:")
            for i, cell_text in enumerate(cells_text):
                print(f"  Cell {i}: '{cell_text}'")

print("\n" + "="*80)
