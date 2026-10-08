# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor

BASE = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/일자별/*** 20260715 Polaris VP preparation ***/별도 제안서 draft"
DARK_BLUE = RGBColor(0x00, 0x04, 0x9C)
BLACK = RGBColor(0x00, 0x00, 0x00)

MARC_EN = ("Additional updates in this document (BV fleet statistics, dry bulk segment age profile, "
    "liquefaction cargo scope, SDARI newbuilding track record, alternative fuel notations, AGS3D "
    "survey technology, and large-ore-carrier terminology in the conclusion) reflect feedback from "
    "Marc Pauchet (Global Market Leader - Dry Bulk), email dated 16 July 2026.")
MARC_KR = ("본 문서의 일부 수치 및 문구(BV 함대 통계, 건식벌크선 구간별 선령 현황, 액상화 대응 화물 범위, "
    "SDARI 신조 실적, 대체연료 노테이션, AGS3D 검사 기술, 결론부 광탄선 관련 표현)는 Marc Pauchet"
    "(Global Market Leader – Dry Bulk)의 2026년 7월 16일자 이메일 피드백을 반영한 것입니다.")

def para(doc, text, bold=False, size=10, color=BLACK):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rpr = r._element.get_or_add_rPr()
    from docx.oxml.ns import qn
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rFonts)
    rFonts.set(qn('w:ascii'), "Arial")
    rFonts.set(qn('w:hAnsi'), "Arial")
    rFonts.set(qn('w:eastAsia'), "맑은 고딕")
    return p

def strip_last_paragraphs(doc, n):
    for _ in range(n):
        p = doc.paragraphs[-1]
        p._p.getparent().remove(p._p)

def build_common_en(doc, project):
    para(doc, "SOURCES — Item-by-Item Detail", bold=True, size=12, color=DARK_BLUE)
    para(doc, "(For internal review only; remove this page before external submission)")
    para(doc, "")
    para(doc, "Executive Summary", bold=True)
    para(doc, "- Background on Vale's 20-vessel triple-fuel bulk carrier newbuilding program and candidate owners: TradeWinds news article (by Alex Gregg-Smith, 2026-07-02)")
    para(doc, "- Detailed analysis of the program and the Qingdao Beihai/Hengli project structure: analysis email from Jeffrey Guo (2026-07-02)")
    para(doc, "- Instruction to draft this Value Proposition for Polaris and the common commercial terms (1st Term FYF Free, Single Class): email from Christophe Vaes to Sang-Jin Han (2026-07-14)")
    para(doc, "- Figures under \"Why Bureau Veritas\": same sources as Section 1.2 below (Dry Bulk Value Proposition, February 2026 edition)")
    para(doc, "")
    para(doc, "Section 1 (BV Support Framework)", bold=True)
    para(doc, "- 1.2 track record figures: Bureau Veritas internal presentation \"Dry Bulk Value Proposition\" (February 2026 edition), Slide 4 \"Almost 200 Years of Experience\", Slide 5 \"A Global Network\", Slide 14 \"BV's Dry Bulk Fleet\", Slide 15 \"Strong within 150-230,000 DWT Ships\", Slide 16 \"Nantong Xiangyu's Activity\", Slide 17 \"Dry Bulk – BV Experience\", Slide 41 \"References\"")
    para(doc, "- 1.3–1.7 (notations, tools, services): based on BV's general product/service literature; applicability to this specific project is [to be updated]")
    para(doc, "- 1.1 dedicated account team: [to be updated – not yet confirmed, no source available]")
    para(doc, "")
    para(doc, "Section 2 (Design Phase)", bold=True)
    para(doc, "- 2.5 liquefaction items: Dry Bulk Value Proposition (February 2026 edition), Slide 19 \"Cargo Liquefaction\", Slide 20 \"Liquefaction is the Biggest Killer\", Slide 21 \"A Clear and Present Danger\", Slide 22 \"Hull Structure\", Slide 23 \"Bureau Veritas Proficiency\", Slide 24 \"Single-Pass Single-Pour\"")
    para(doc, "- 2.2 premium project organization: Slide 8 \"Pre-Phase Project Management\", Slide 9 \"Project Management\"")
    para(doc, "- 2.6 Qingdao office coverage: Slide 7 \"BV Present in Every Main S/B Location in China\"")
    para(doc, "- 2.1, 2.3, 2.4, 2.7, 2.8: based on BV's general rules/service literature; project-specific details are [to be updated]")
    para(doc, "")
    para(doc, "Section 3 (Construction Phase)", bold=True)
    if project == "P2":
        para(doc, "- 3.4 Hengli site organization and PPM: email from Mustafa Bilgin to Christophe Vaes and others (2026-07-09), attachment \"Hengli SY - 2 x Newcastlemax BC for Polaris.pptx\"")
    para(doc, "- 3.1, 3.2, 3.3: based on BV's general Premium Project Management literature; dedicated personnel/organization for this project are [to be updated]")
    para(doc, "")
    para(doc, "Section 4 (Project-Specific Proposal)", bold=True)
    para(doc, "- Project background: email from Christophe Vaes to Chris Greenwood and others (2026-07-07, re: Polaris Busan meeting)")
    if project == "P2":
        para(doc, "- Hengli design-related background details (ABS drawing approval status, etc.): email from Cindy Cai (ZI XU CAI) to Christophe Vaes, Mustafa Bilgin and others (2026-07-10) — details unsuitable for external disclosure (e.g. competitor information) have been excluded from this document")
    para(doc, "")
    para(doc, "Section 5 (Commercial Terms)", bold=True)
    para(doc, "- Email from Christophe Vaes to Sang-Jin Han (2026-07-14)")
    para(doc, "")
    para(doc, "Sections 6 and 7", bold=True)
    para(doc, "- Reflects the drafting team's (Bureau Veritas Korea) own recommendations; no external or internal source document applies.")
    para(doc, "")
    para(doc, "Section 8 (List of Annexes)", bold=True)
    para(doc, "- Annex A: refers to the standalone file (original source: Dry Bulk Value Proposition, February 2026 edition)")
    if project == "P1":
        para(doc, "- Annex C, D: extracted/organized from the Annex A source material and Qingdao Beihai project-specific data; content marked [to be updated] where not yet confirmed")
    else:
        para(doc, "- Annex B: attachment to email from Mustafa Bilgin to Christophe Vaes and others (2026-07-09), \"Hengli SY - 2 x Newcastlemax BC for Polaris.pptx\"")
    para(doc, "- Annex E: project organization chart and team CVs; personnel to be confirmed upon contract award [to be updated]")
    para(doc, "- Annex F: commercial terms per email from Christophe Vaes to Sang-Jin Han (2026-07-14)")
    para(doc, "")
    para(doc, "Note: The email sources above refer to Bureau Veritas Korea's internal mailbox and can be cross-checked against the original emails if required.")
    para(doc, "")
    para(doc, MARC_EN)

def build_common_kr(doc, project):
    para(doc, "출처 (Sources) — 항목별 상세 출처", bold=True, size=12, color=DARK_BLUE)
    para(doc, "(내부 검토용 – 외부 제출본에서는 본 페이지 제외 요망)")
    para(doc, "")
    para(doc, "Executive Summary", bold=True)
    para(doc, "- Vale의 20척 규모 triple-fuel 대형 벌크선 신조 프로그램 및 후보 선주 관련 동향: TradeWinds 뉴스 기사 (Alex Gregg-Smith 작성, 2026-07-02)")
    para(doc, "- 프로그램 세부 분석 및 Qingdao Beihai/Hengli 프로젝트 구도 파악: Jeffrey Guo 분석 이메일 (2026-07-02)")
    para(doc, "- Polaris 대상 Value Proposition 작성 지시 및 공통 상업 조건(1st Term FYF Free, 단일 클래스) 전제: Christophe Vaes → Sang-Jin Han 이메일 (2026-07-14)")
    para(doc, "- \"왜 Bureau Veritas인가\" 항목의 실적 수치: 아래 1.2 항목 출처와 동일 (Dry Bulk Value Proposition, 2026년 2월판)")
    para(doc, "")
    para(doc, "1장 (BV 지원 체계)", bold=True)
    para(doc, "- 1.2 실적 수치: Bureau Veritas 내부 자료 \"Dry Bulk Value Proposition\"(2026년 2월판) Slide 4 \"Almost 200 Years of Experience\", Slide 5 \"A Global Network\", Slide 14 \"BV's Dry Bulk Fleet\", Slide 15 \"Strong within 150-230,000 DWT Ships\", Slide 16 \"Nantong Xiangyu's Activity\", Slide 17 \"Dry Bulk – BV Experience\", Slide 41 \"References\"")
    para(doc, "- 1.3~1.7 항목 중 노테이션·툴·서비스 설명: BV 일반 제품/서비스 소개 자료 기반 서술 (본 프로젝트 특정 적용 여부는 [to be updated])")
    para(doc, "- 1.1 전담 조직: [to be updated – 담당자 미확정으로 출처 없음]")
    para(doc, "")
    para(doc, "2장 (설계 단계)", bold=True)
    para(doc, "- 2.5 액상화 대응 항목: Dry Bulk Value Proposition(2026년 2월판) Slide 19 \"Cargo Liquefaction\", Slide 20 \"Liquefaction is the Biggest Killer\", Slide 21 \"A Clear and Present Danger\", Slide 22 \"Hull Structure\", Slide 23 \"Bureau Veritas Proficiency\", Slide 24 \"Single-Pass Single-Pour\"")
    para(doc, "- 2.2 프리미엄 프로젝트 조직 관련 일반 서술: Slide 8 \"Pre-Phase Project Management\", Slide 9 \"Project Management\"")
    para(doc, "- 2.6 칭다오 사업소 커버리지: Slide 7 \"BV Present in Every Main S/B Location in China\"")
    para(doc, "- 2.1, 2.3, 2.4, 2.7, 2.8 항목: BV 일반 규정/서비스 소개 자료 기반 서술이며, 본 프로젝트 특정 적용 세부사항은 [to be updated]")
    para(doc, "")
    para(doc, "3장 (건조 단계)", bold=True)
    if project == "P2":
        para(doc, "- 3.4 Hengli 현장 조직 및 PPM: Mustafa Bilgin → Christophe Vaes 등 이메일 (2026-07-09), 첨부 \"Hengli SY - 2 x Newcastlemax BC for Polaris.pptx\"")
    para(doc, "- 3.1, 3.2, 3.3 항목: BV 일반 프리미엄 프로젝트 매니지먼트 체계 소개 자료 기반 서술이며, 본 프로젝트 전담 인력/조직은 [to be updated]")
    para(doc, "")
    para(doc, "4장 (프로젝트별 제안)", bold=True)
    para(doc, "- 프로젝트 개요 및 배경: Christophe Vaes → Chris Greenwood 등 이메일 (2026-07-07, Polaris 부산 미팅 관련)")
    if project == "P2":
        para(doc, "- Hengli 설계 관련 세부사항(ABS 도면 승인 현황 등 내부 배경 정보): Cindy Cai(ZI XU CAI) → Christophe Vaes, Mustafa Bilgin 등 이메일 (2026-07-10) — 본 문서에는 대외 제출에 부적합한 세부사항(경쟁사 정보 등)은 제외하고 반영")
    para(doc, "")
    para(doc, "5장 (상업 조건)", bold=True)
    para(doc, "- Christophe Vaes → Sang-Jin Han 이메일 (2026-07-14)")
    para(doc, "")
    para(doc, "6장, 7장", bold=True)
    para(doc, "- 작성자(Bureau Veritas Korea) 판단에 따른 서술이며, 별도의 외부/내부 자료 출처는 없습니다.")
    para(doc, "")
    para(doc, "8장 (Annex 목록)", bold=True)
    para(doc, "- Annex A: 자체 문서 참조 (원출처: Dry Bulk Value Proposition, 2026년 2월판)")
    if project == "P1":
        para(doc, "- Annex C, D: Annex A 원자료 및 Qingdao Beihai 프로젝트 관련 자료에서 발췌·정리; 미확정 항목은 [to be updated]로 표기")
    else:
        para(doc, "- Annex B: Mustafa Bilgin 이메일 (2026-07-09) 첨부파일 \"Hengli SY - 2 x Newcastlemax BC for Polaris.pptx\"")
    para(doc, "- Annex E: 프로젝트 조직도 및 팀 CV; 담당 인력은 계약 진행 시 확정 예정 [to be updated]")
    para(doc, "- Annex F: 상업 조건, Christophe Vaes → Sang-Jin Han 이메일 (2026-07-14) 기준")
    para(doc, "")
    para(doc, "참고: 상기 이메일 출처는 Bureau Veritas Korea 내부 이메일함 기준이며, 필요 시 원본 이메일 열람을 통해 교차 확인 가능합니다.")
    para(doc, "")
    para(doc, MARC_KR)

import os

TARGETS = [
    ("Project 1 (Qingdao VLOC)/Value Proposition to Polaris - Project 1 (Qingdao VLOC) (EN).docx", "P1", "EN"),
    ("Project 1 (Qingdao VLOC)/Value Proposition to Polaris - Project 1 (Qingdao VLOC) (KR).docx", "P1", "KR"),
    ("Project 2 (Hengli BC)/Value Proposition to Polaris - Project 2 (Hengli BC) (EN).docx", "P2", "EN"),
    ("Project 2 (Hengli BC)/Value Proposition to Polaris - Project 2 (Hengli BC) (KR).docx", "P2", "KR"),
]

for rel, project, lang in TARGETS:
    path = os.path.join(BASE, rel)
    doc = docx.Document(path)
    # remove the minimal "Sources"/"출처" heading + note paragraphs I added previously (last 2 paragraphs)
    last_text = doc.paragraphs[-1].text
    second_last = doc.paragraphs[-2].text
    if ("Marc Pauchet" in last_text or "Marc Pauchet" in second_last):
        strip_last_paragraphs(doc, 2)
    else:
        print("WARNING: unexpected tail, skipping strip for", rel)

    if lang == "EN":
        build_common_en(doc, project)
    else:
        build_common_kr(doc, project)

    doc.save(path)
    print("done:", rel)

print("ALL DONE")
