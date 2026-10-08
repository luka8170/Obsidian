#!/usr/bin/env python3

from docx import Document
import sys

print("="*80)
print("ANNEX C 검사 횟수 상세 비교")
print("="*80)

# 1. Read Word file - FYF Agreement (내가 보낸 파일)
print("\n[1] 내가 보낸 FYF Agreement Word 파일 읽기...")
word_file = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/내부 회의 내부 작업/**2. QG 계약서/POSSM/최종 정리/내가 보낸 파일들/word/25.8.6. 송부) FYF AGREEMENT - POSSM(PAN OCEAN)_(as of Apr 30 2025) rev2.docx"

try:
    doc = Document(word_file)
    
    print(f"Total tables: {len(doc.tables)}")
    
    word_surveys = {}
    
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            
            if len(cells) >= 3:
                survey_code = cells[0]
                survey_name = cells[1] if len(cells) > 1 else ""
                survey_count = cells[2] if len(cells) > 2 else ""
                
                # Look for specific surveys
                if survey_code in ["DOK", "TPSC", "TSBC"]:
                    word_surveys[survey_code] = {
                        "name": survey_name,
                        "count": survey_count,
                        "table_idx": table_idx,
                        "row_idx": row_idx
                    }
                    print(f"  Found {survey_code}: {survey_name} = '{survey_count}' (Table {table_idx}, Row {row_idx})")
    
    print(f"\n내가 보낸 파일 검사 횟수:")
    print(f"  DOK: {word_surveys.get('DOK', {}).get('count', 'NOT FOUND')}")
    print(f"  TPSC: {word_surveys.get('TPSC', {}).get('count', 'NOT FOUND')}")
    print(f"  TSBC: {word_surveys.get('TSBC', {}).get('count', 'NOT FOUND')}")

except Exception as e:
    print(f"Error reading Word file: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "="*80)
print("파일 읽기 완료")
print("="*80)
