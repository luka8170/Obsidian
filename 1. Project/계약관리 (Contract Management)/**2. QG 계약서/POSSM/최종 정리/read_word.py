#!/usr/bin/env python3
from docx import Document

print("="*100)
print("WORD 파일 분석 - ANNEX C 검사 횟수")
print("="*100)

word_file = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/내부 회의 내부 작업/**2. QG 계약서/POSSM/최종 정리/내가 보낸 파일들/word/25.8.6. 송부) FYF AGREEMENT - POSSM(PAN OCEAN)_(as of Apr 30 2025) rev2.docx"

doc = Document(word_file)
print(f"\nTotal tables in document: {len(doc.tables)}\n")

# Search all tables for survey data
dok_count = None
tpsc_count = None
tsbc_count = None

for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        row_text = "|".join(cells)
        
        # Look for DOK
        if "DOK" in cells[0] if len(cells) > 0 else False:
            if "Periodical Bottom" in row_text or "Dry Dock" in row_text:
                if len(cells) >= 3:
                    dok_count = cells[2]
                    print(f"✓ Found DOK in Table {table_idx}, Row {row_idx}")
                    print(f"  Full row: {cells}")
                    print(f"  Count: {dok_count}\n")
        
        # Look for TPSC
        if "TPSC" in cells[0] if len(cells) > 0 else False:
            if "Portside" in row_text or "Tailshaft" in row_text:
                if len(cells) >= 3:
                    tpsc_count = cells[2]
                    print(f"✓ Found TPSC in Table {table_idx}, Row {row_idx}")
                    print(f"  Full row: {cells}")
                    print(f"  Count: {tpsc_count}\n")
        
        # Look for TSBC
        if "TSBC" in cells[0] if len(cells) > 0 else False:
            if "Starboard" in row_text or "Tailshaft" in row_text:
                if len(cells) >= 3:
                    tsbc_count = cells[2]
                    print(f"✓ Found TSBC in Table {table_idx}, Row {row_idx}")
                    print(f"  Full row: {cells}")
                    print(f"  Count: {tsbc_count}\n")

print("="*100)
print("내가 보낸 Word 파일 검사 횟수 (ANNEX C)")
print("="*100)
print(f"DOK  (Dry Dock Survey):        {dok_count if dok_count else 'NOT FOUND'}")
print(f"TPSC (Portside Tailshaft):     {tpsc_count if tpsc_count else 'NOT FOUND'}")
print(f"TSBC (Starboard Tailshaft):    {tsbc_count if tsbc_count else 'NOT FOUND'}")
print("="*100)
