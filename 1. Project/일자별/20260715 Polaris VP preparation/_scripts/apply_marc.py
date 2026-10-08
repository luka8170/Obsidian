import docx, os, copy
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE = "/Users/sangjinhan/Library/Mobile Documents/iCloud~md~obsidian/Documents/Obsidian/1. Project/일자별/*** 20260715 Polaris VP preparation ***/별도 제안서 draft"

DARK_BLUE = RGBColor(0x00, 0x04, 0x9C)
BLACK = RGBColor(0x00, 0x00, 0x00)

MAIN_FILES = [
    "Project 1 (Qingdao VLOC)/Value Proposition to Polaris - Project 1 (Qingdao VLOC) (EN).docx",
    "Project 1 (Qingdao VLOC)/Value Proposition to Polaris - Project 1 (Qingdao VLOC) (KR).docx",
    "Project 2 (Hengli BC)/Value Proposition to Polaris - Project 2 (Hengli BC) (EN).docx",
    "Project 2 (Hengli BC)/Value Proposition to Polaris - Project 2 (Hengli BC) (KR).docx",
]
ANNEXA_FILES = [
    "Project 1 (Qingdao VLOC)/Annexes/Annex A - BV Dry Bulk VLOC Track Record (EN).docx",
    "Project 1 (Qingdao VLOC)/Annexes/Annex A - BV Dry Bulk VLOC Track Record (KR).docx",
    "Project 2 (Hengli BC)/Annexes/Annex A - BV Dry Bulk VLOC Track Record (EN).docx",
    "Project 2 (Hengli BC)/Annexes/Annex A - BV Dry Bulk VLOC Track Record (KR).docx",
]
ANNEXC_FILES = [
    "Project 1 (Qingdao VLOC)/Annexes/Annex C - BV VLOC Liquefaction Assessment (EN).docx",
    "Project 1 (Qingdao VLOC)/Annexes/Annex C - BV VLOC Liquefaction Assessment (KR).docx",
]

# (old, new) substring replacement pairs. Applied wherever found; no-op otherwise.
REPLACEMENTS = [
    # --- fleet stats ---
    ("1,213 vessels (approx. 49.2 million GT)", "1,203 vessels (approx. 50.4 million GT)"),
    ("1,213 vessels / 49.2 million GT", "1,203 vessels / 50.4 million GT"),
    ("1,213척(약 49.2백만 GT)", "1,203척(약 50.4백만 GT)"),
    ("1,213척/49.2백만 GT", "1,203척/50.4백만 GT"),
    ("average age of 9 years", "average age of 12.0 years"),
    ("average age 9 years", "average age 12.0 years"),
    ("평균 선령 9년", "평균 선령 12.0년"),
    ("IACS average (11.8 years)", "IACS average (12.6 years)"),
    ("IACS average of 11.8 years", "IACS average of 12.6 years"),
    ("IACS 평균(11.8년)", "IACS 평균(12.6년)"),
    ("IACS 평균 11.8년", "IACS 평균 12.6년"),
    ("dry bulk fleet average (12.8 years)", "dry bulk fleet average (13.0 years)"),
    ("dry bulk average of 12.8 years", "dry bulk average of 13.0 years"),
    ("건식벌크선 평균(12.8년)", "건식벌크선 평균(13.0년)"),
    ("건식벌크선 평균 12.8년", "건식벌크선 평균 13.0년"),
    ("전체 건식 벌크선 평균(12.8년)", "전체 건식 벌크선 평균(13.0년)"),
    ("41 newbuilding deliveries in the past 12 months",
     "76 newbuilding deliveries (approx. 4.8 million GT) in the past 12 months"),
    ("41 new bulk carriers were delivered under BV class in the past 12 months.",
     "76 new bulk carriers (approx. 4.8 million GT) were delivered under BV class in the past 12 months."),
    ("41 deliveries in the past 12 months", "76 deliveries (approx. 4.8 million GT) in the past 12 months"),
    ("최근 12개월간 신규 인도 41척", "최근 12개월간 신규 인도 76척(약 4.8백만 GT)"),
    ("최근 12개월간 41척의 신규 벌크선이", "최근 12개월간 76척(약 4.8백만 GT)의 신규 벌크선이"),
    ("최근 12개월 41척 인도", "최근 12개월 76척(약 4.8백만 GT) 인도"),

    # --- segment age claim ---
    ("where all vessels under BV class are less than 5 years old",
     "where vessels under BV class average less than 15 years old"),
    ("해당 구간 BV 선급선은 전량 선령 5년 미만", "해당 구간 BV 선급선은 평균 선령 15년 미만"),

    # --- liquefaction commodity narrative (main VP bullet + prose) ---
    ("Dedicated research and a proven safety record addressing the liquefaction risk of iron ore fines, Vale's principal cargo, including a VLOC liquefaction stability assessment",
     "Dedicated research and a proven safety record addressing the liquefaction risk of iron ore fines, manganese ore fines and bauxite fines. The former two commodities are among Vale's main cargos while bauxite is the fastest growing commodity for large dry bulk ships. Bureau Veritas' track record includes VLOC liquefaction stability assessment."),
    ("Vale의 주력 화물인 철광석(Iron Ore Fines)의 액상화(Liquefaction) 리스크에 대한 전용 연구 및 VLOC 액상화 안정성 평가 실적 보유",
     "철광석(Iron Ore Fines), 망간광 분광석(Manganese Ore Fines), 보크사이트 분광석(Bauxite Fines)의 액상화(Liquefaction) 리스크에 대한 전용 연구 실적 보유 – 앞의 두 화물은 Vale의 주요 화물이며, 보크사이트는 대형 건식벌크선 화물 중 가장 빠르게 성장하는 화물임. VLOC 액상화 안정성 평가 실적 포함"),

    # --- SDARI / Nantong Xiangyu stat enrichment ---
    ("55% of SDARI-design bulk carriers built or on order at Nantong Xiangyu Shipyard are under BV class (24% of that shipyard's total volume)",
     "55% of SDARI-design bulk carriers built or on order at Nantong Xiangyu Shipyard are under BV class (24% of that shipyard's total volume); BV also oversees 44 SDARI-designed vessels currently under construction across China, representing 14% of all SDARI-designed dry bulk vessels"),
    ("(24% of the shipyard's total volume). Further references include",
     "(24% of the shipyard's total volume). BV currently oversees 44 SDARI-designed vessels under construction across China, representing 14% of all SDARI-designed dry bulk vessels. Further references include"),
    ("Nantong Xiangyu 조선소 SDARI 설계 벌크선의 55%가 BV 선급 (전체 건조·발주 물량의 24%)",
     "Nantong Xiangyu 조선소 SDARI 설계 벌크선의 55%가 BV 선급 (전체 건조·발주 물량의 24%); 이 외에도 BV는 현재 중국 전역에서 건조 중인 SDARI 설계 선박 44척을 관리하고 있으며, 이는 전체 SDARI 설계 건식벌크선의 14%에 해당함"),
    ("Nantong Xiangyu 조선소에서 건조·발주된 SDARI 설계 벌크선의 55%가 BV 선급(조선소 전체 물량 기준 24%)입니다. 참고 실적으로",
     "Nantong Xiangyu 조선소에서 건조·발주된 SDARI 설계 벌크선의 55%가 BV 선급(조선소 전체 물량 기준 24%)입니다. 이 외에도 BV는 현재 중국 전역에서 건조 중인 SDARI 설계 선박 44척을 관리 중이며, 이는 전체 SDARI 설계 건식벌크선의 14%에 해당합니다. 참고 실적으로"),
    ("Nantong Xiangyu Shipyard: 24% of all dry bulk vessels built or on order there are under BV class (55% among SDARI-designed vessels)",
     "Nantong Xiangyu Shipyard: 24% of all dry bulk vessels built or on order there are under BV class (55% among SDARI-designed vessels); BV also oversees 44 SDARI-designed vessels under construction across China (14% of all SDARI-designed dry bulk vessels)"),
    ("Nantong Xiangyu 조선소: 건조/건조 예정 건식벌크선의 24%가 BV 선급 (SDARI 설계 선박 중에서는 55%가 BV 선급)",
     "Nantong Xiangyu 조선소: 건조/건조 예정 건식벌크선의 24%가 BV 선급 (SDARI 설계 선박 중에서는 55%가 BV 선급); 이 외 중국 전역에서 건조 중인 SDARI 설계 선박 44척을 BV가 관리 중이며, 이는 전체 SDARI 설계 건식벌크선의 14%에 해당"),

    # --- ethanol addition ---
    ("methanol, ammonia, LNG, etc.", "methanol, ethanol, ammonia, LNG, etc."),
    ("methanol, ammonia, LNG 등", "methanol, ethanol, ammonia, LNG 등"),

    # --- conclusion wording (VLOC -> large ore carrier framing) ---
    ("Bureau Veritas confirms its readiness to support Polaris on Vale's large bulk carrier/VLOC newbuilding program. Drawing on close to 200 years of classification history, a deep track record in the dry bulk/VLOC segment,",
     "Bureau Veritas confirms its readiness to support Polaris on Vale's large bulk and ore carrier newbuilding program. Drawing on close to 200 years of classification history, a deep track record in the dry bulk and ore carrier segment,"),
    ("Vale의 대형 벌크선/VLOC 신조 프로그램에서 Polaris와 함께할 준비가 되어 있음을 말씀드립니다. 200년에 가까운 선급 역사, 건식벌크선/VLOC 분야의 깊이 있는 실적,",
     "Vale의 대형 벌크선 및 광탄선(ore carrier) 신조 프로그램에서 Polaris와 함께할 준비가 되어 있음을 말씀드립니다. 200년에 가까운 선급 역사, 건식벌크선 및 광탄선 분야의 깊이 있는 실적,"),

    # --- Annex A title / reference wording ---
    ("Bureau Veritas Dry Bulk / VLOC Capability and Track Record",
     "Bureau Veritas Dry Bulk / Large Ore Carrier Capability and Track Record"),
    ("Bureau Veritas 건식벌크선 / VLOC 부문 역량 및 실적 자료",
     "Bureau Veritas 건식벌크선 / 대형 광탄선(Large Ore Carrier) 부문 역량 및 실적 자료"),
]

def find_para(doc, exact_text):
    for p in doc.paragraphs:
        if p.text.strip() == exact_text.strip():
            return p
    return None

def insert_after(anchor_p, new_text):
    """Duplicate anchor_p's XML (preserves style/formatting) and insert right after it,
    then overwrite the text of the duplicate with new_text."""
    new_p_elem = copy.deepcopy(anchor_p._p)
    anchor_p._p.addnext(new_p_elem)
    new_para = docx.text.paragraph.Paragraph(new_p_elem, anchor_p._parent)
    runs = new_para.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ''
    else:
        new_para.add_run(new_text)
    return new_para

def replace_in_doc(doc, pairs):
    changed = 0
    for p in doc.paragraphs:
        full = p.text
        new_full = full
        for old, new in pairs:
            if old in new_full:
                new_full = new_full.replace(old, new)
        if new_full != full:
            # apply the change per-run when possible; fallback to whole-paragraph rewrite
            applied = False
            for old, new in pairs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        applied = True
            # verify; if still stale (phrase spanned multiple runs), rebuild paragraph text
            if p.text != new_full:
                runs = p.runs
                if runs:
                    runs[0].text = new_full
                    for r in runs[1:]:
                        r.text = ''
                else:
                    p.add_run(new_full)
            changed += 1
    return changed

AGS3D_ANCHOR_EN = "- Machinery Maintenance Portal / VeriSTAR Equipment: digital portals for machinery maintenance tracking and equipment certificate management"
AGS3D_TEXT_EN = "- AGS3D: Bureau Veritas' drone-based survey and inspection technology, with deployment planned in Brazil"
AGS3D_ANCHOR_KR = "- Machinery Maintenance Portal / VeriSTAR Equipment: 기관 유지보수 및 기자재 인증서 관리를 위한 디지털 포털"
AGS3D_TEXT_KR = "- AGS3D: Bureau Veritas의 드론 기반 검사 기술 – 브라질 배치 예정"

ANNEXC_ANCHOR_EN = "- SINGLEPASSLOADING / GRABLOADING notations"
ANNEXC_TEXT_EN = "- Liquefaction coverage also extends to manganese ore fines and bauxite fines, in addition to iron ore fines"
ANNEXC_ANCHOR_KR = "- SINGLEPASSLOADING / GRABLOADING 노테이션 관련 내용"
ANNEXC_TEXT_KR = "- 액상화 대응 범위는 철광석 분광석 외에 망간광 분광석, 보크사이트 분광석까지 포함"

ANNEXA_HIST_ANCHOR_EN = "This Annex presents a curated extract from Bureau Veritas' internal \"Dry Bulk Value Proposition\" (February 2026 edition), limited to the items directly relevant to this project (Vale's 210,000 DWT VLOC/Bulk Carrier newbuilding program). It is not the full presentation, but a selection of the core track record and capability data relevant to the Polaris proposal."
ANNEXA_HIST_TEXT_EN = "For additional historical context, cross-referencing BV's Veristar and Neptune/Tableau records identifies 8 ships over 220,000 DWT classed by BV over time, the largest being the 275,726 DWT MAIN ORE (built at Uljanik, Croatia). Four of these vessels were disposed of while still under BV class; the other four changed class and have since been demolished. [Source: Marc Pauchet feedback, 16 Jul 2026]"
ANNEXA_HIST_ANCHOR_KR = "본 Annex는 Bureau Veritas 내부 자료인 \"Dry Bulk Value Proposition\"(2026년 2월판) 중 본 프로젝트(Vale向け 210,000 DWT VLOC/Bulk Carrier 신조)와 직접 관련된 항목만 선별하여 정리한 자료입니다. 전체 프레젠테이션 자료가 아닌, Polaris 제안에 필요한 핵심 실적/역량 데이터만 발췌하였습니다."
ANNEXA_HIST_TEXT_KR = "추가 참고 사항으로, BV의 Veristar 및 Neptune/Tableau 기록을 교차 확인한 결과 과거 BV가 선급한 220,000 DWT 이상 선박이 8척 확인되었으며, 이 중 가장 큰 선박은 275,726 DWT급 MAIN ORE(Uljanik, Croatia 건조)입니다. 이 중 4척은 BV 선급 상태에서 매각되었고, 나머지 4척은 타 선급으로 이전 후 해체되었습니다. [출처: Marc Pauchet 피드백, 2026-07-16]"

def is_kr(path):
    return path.rstrip(".docx").endswith("(KR)")

MARC_NOTE_EN_MAIN = ("Sources", [
    "Content updates in this document (BV fleet statistics, dry bulk segment age profile, "
    "liquefaction cargo scope, SDARI newbuilding track record, alternative fuel notations, "
    "AGS3D survey technology, and large-ore-carrier terminology in the conclusion) reflect "
    "feedback from Marc Pauchet (Global Market Leader - Dry Bulk), email dated 16 July 2026."
])
MARC_NOTE_KR_MAIN = ("출처", [
    "본 문서의 일부 수치 및 문구(BV 함대 통계, 건식벌크선 구간별 선령 현황, 액상화 대응 화물 범위, "
    "SDARI 신조 실적, 대체연료 노테이션, AGS3D 검사 기술, 결론부 광탄선 관련 표현)는 "
    "Marc Pauchet(Global Market Leader – Dry Bulk)의 2026년 7월 16일자 이메일 피드백을 반영한 것입니다."
])

ANNEXA_SOURCE_LINE_EN = ("- Fleet statistics (1,203 vessels / 50.4 million GT, average age 12.0 years), "
    "150,000-230,000 DWT segment age profile, SDARI 44-vessel/14% newbuilding stat, and the "
    ">220,000 DWT historical fleet reference: Marc Pauchet feedback, email dated 16 July 2026")
ANNEXA_SOURCE_LINE_KR = ("- 함대 통계(1,203척/50.4백만 GT, 평균 선령 12.0년), 150,000-230,000 DWT 구간 선령 현황, "
    "SDARI 44척/14% 신조 실적, 220,000 DWT 이상 과거 선박 이력: Marc Pauchet 피드백, 2026년 7월 16일자 이메일")

ANNEXC_SOURCE_LINE_EN = "Additional liquefaction cargo scope (manganese ore fines, bauxite fines) added per Marc Pauchet feedback, email dated 16 July 2026."
ANNEXC_SOURCE_LINE_KR = "액상화 대응 화물 범위(망간광 분광석, 보크사이트 분광석) 추가 반영 – Marc Pauchet 피드백, 2026년 7월 16일자 이메일."

results = []

def process(rel, kind):
    path = os.path.join(BASE, rel)
    doc = docx.Document(path)
    kr = is_kr(rel)
    n_repl = replace_in_doc(doc, REPLACEMENTS)
    inserted = []

    if kind == "main":
        anchor = find_para(doc, AGS3D_ANCHOR_KR if kr else AGS3D_ANCHOR_EN)
        if anchor:
            insert_after(anchor, AGS3D_TEXT_KR if kr else AGS3D_TEXT_EN)
            inserted.append("AGS3D")
    elif kind == "annexA":
        anchor = find_para(doc, ANNEXA_HIST_ANCHOR_KR if kr else ANNEXA_HIST_ANCHOR_EN)
        if anchor:
            insert_after(anchor, ANNEXA_HIST_TEXT_KR if kr else ANNEXA_HIST_TEXT_EN)
            inserted.append("HistFleet")
    elif kind == "annexC":
        anchor = find_para(doc, ANNEXC_ANCHOR_KR if kr else ANNEXC_ANCHOR_EN)
        if anchor:
            insert_after(anchor, ANNEXC_TEXT_KR if kr else ANNEXC_TEXT_EN)
            inserted.append("ManganeseBauxite")

    changed = n_repl > 0 or len(inserted) > 0

    if changed:
        if kind == "main":
            heading, lines = MARC_NOTE_KR_MAIN if kr else MARC_NOTE_EN_MAIN
            doc.add_paragraph("")
            h = doc.add_paragraph(heading)
            for r in h.runs:
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = DARK_BLUE
            for line in lines:
                b = doc.add_paragraph(line)
                for r in b.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = BLACK
        elif kind == "annexA":
            src_heading = None
            for p in doc.paragraphs:
                if p.text.strip().upper().startswith("SOURCES"):
                    src_heading = p
                    break
            line = ANNEXA_SOURCE_LINE_KR if kr else ANNEXA_SOURCE_LINE_EN
            if src_heading is not None:
                insert_after(src_heading, line)
            else:
                b = doc.add_paragraph(line)
        elif kind == "annexC":
            line = ANNEXC_SOURCE_LINE_KR if kr else ANNEXC_SOURCE_LINE_EN
            doc.add_paragraph(line)

    doc.save(path)
    results.append((rel, n_repl, inserted, changed))

for f in MAIN_FILES:
    process(f, "main")
for f in ANNEXA_FILES:
    process(f, "annexA")
for f in ANNEXC_FILES:
    process(f, "annexC")

print("\n=== SUMMARY ===")
for rel, n_repl, inserted, changed in results:
    print(f"changed={str(changed):5s} paras_replaced={n_repl:2d} inserted={inserted} | {rel}")
